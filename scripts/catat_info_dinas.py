#!/usr/bin/env python3
"""
catat_info_dinas.py — Append satu baris info dari Dinas Pendidikan / Kemendikdasmen
ke spreadsheet Excel (.xlsx) lokal.

Cara pakai (dari Hermes / command line):
  echo '{"isi":"Surat edaran UN mendatang","sumber":"Kemendikdasmen","kategori":"Surat Edaran"}' \
    | python scripts/catat_info_dinas.py

Input: JSON di stdin dengan field opsional:
  isi            (wajib) — isi info surat/pengumuman
  tanggal_info   — tanggal tertulis di surat (kosong = "-")
  sumber         — default "Dinas/Kemendikdasmen"
  kategori       — Surat Edaran | Juknis | Lomba | Kalender | Lainnya (default Lainnya)
  tindak_lanjut  — rencana sekolah (default "-")
  status         — Belum ditindak | Proses | Selesai (default Belum ditindak)
  pj             — penanggung jawab (default "-")
  xlsx           — override path file (default env INFO_DINAS_XLSX atau D:/2026-2027/info_dinas.xlsx)

Output: JSON {"ok":true,"no":N,"path":"...","tanggal_catat":"..."}

Bila stdin bukan JSON, seluruh teks dianggap field `isi`.
Path opsional disimpan di env INFO_DINAS_XLSX agar portable antar sekolah.
"""
import sys
import os
import json
import datetime

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print(json.dumps({"ok": False, "error": "openpyxl tidak terpasang. Install: pip install openpyxl"}))
    sys.exit(2)

DEFAULT_PATH = os.environ.get("INFO_DINAS_XLSX", r"D:/2026-2027/info_dinas.xlsx")
HEADERS = ["No", "Tanggal Catat", "Tanggal Info", "Sumber", "Kategori",
           "Isi Info", "Tindak Lanjut", "Status", "PJ"]


def _wib_now():
    # WIB = UTC+7
    return datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=7)))


def _append_row(path, data):
    """Append satu baris ke xlsx. Return (no, tanggal_catat)."""
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        parent = os.path.dirname(path)
        if parent:
            os.makedirs(parent, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)

    data_rows = max(ws.max_row - 1, 0)
    no = data_rows + 1
    tcatat = _wib_now().strftime("%Y-%m-%d %H:%M")

    row = [
        no,
        tcatat,
        data.get("tanggal_info") or "-",
        data.get("sumber") or "Dinas/Kemendikdasmen",
        data.get("kategori") or "Lainnya",
        data.get("isi") or "",
        data.get("tindak_lanjut") or "-",
        data.get("status") or "Belum ditindak",
        data.get("pj") or "-",
    ]
    ws.append(row)
    wb.save(path)
    return no, tcatat


def main():
    raw = sys.stdin.read().strip()
    if not raw:
        print(json.dumps({"ok": False, "error": "stdin kosong"}))
        sys.exit(1)

    try:
        data = json.loads(raw)
        if not isinstance(data, dict):
            data = {"isi": str(data)}
    except Exception:
        # Bukan JSON -> anggap seluruh teks adalah isi info
        data = {"isi": raw}

    isi = data.get("isi", "").strip()
    if not isi:
        print(json.dumps({"ok": False, "error": "field 'isi' wajib diisi dan tidak boleh kosong"}))
        sys.exit(1)

    path = data.get("xlsx") or DEFAULT_PATH

    try:
        no, tcatat = _append_row(path, data)
    except PermissionError:
        print(json.dumps({"ok": False, "error": f"file terkunci/sedang dibuka: {path}"}))
        sys.exit(3)
    except Exception as e:
        print(json.dumps({"ok": False, "error": f"gagal tulis xlsx: {e}"}))
        sys.exit(4)

    print(json.dumps({
        "ok": True,
        "no": no,
        "path": path,
        "tanggal_catat": tcatat,
    }, ensure_ascii=False))


# -----------------------------------------------------------------------
# Self-test: jalankan dengan --self-check
# -----------------------------------------------------------------------
def _self_check():
    """Verifikasi: append 2 baris ke file temp, cek header + baris + nomor urut."""
    import tempfile
    import subprocess

    tmp = os.path.join(tempfile.gettempdir(), "_edudev_selfcheck.xlsx")
    if os.path.exists(tmp):
        os.remove(tmp)

    # Baris 1
    p = subprocess.run(
        [sys.executable, os.path.abspath(__file__)],
        input=json.dumps({"isi": "test baris 1", "xlsx": tmp}, ensure_ascii=False),
        capture_output=True, text=True,
    )
    r1 = json.loads(p.stdout)
    assert r1["ok"] is True, f"baris 1 gagal: {r1}"
    assert r1["no"] == 1, f"no baris 1 salah: {r1['no']}"

    # Baris 2
    p = subprocess.run(
        [sys.executable, os.path.abspath(__file__)],
        input=json.dumps({"isi": "test baris 2", "xlsx": tmp}, ensure_ascii=False),
        capture_output=True, text=True,
    )
    r2 = json.loads(p.stdout)
    assert r2["ok"] is True, f"baris 2 gagal: {r2}"
    assert r2["no"] == 2, f"no baris 2 salah: {r2['no']}"

    # Verifikasi isi file
    wb = load_workbook(tmp)
    ws = wb.active
    assert ws.max_row == 3, f"baris seharusnya 3 (header+2 data), dapat {ws.max_row}"
    assert ws.cell(1, 1).value == "No", "header kolom 1 bukan 'No'"
    assert ws.cell(2, 6).value == "test baris 1", f"isi baris 2 kolom 6 salah: {ws.cell(2,6).value}"
    assert ws.cell(3, 6).value == "test baris 2", f"isi baris 3 kolom 6 salah: {ws.cell(3,6).value}"

    os.remove(tmp)
    print("SELF-CHECK PASSED: 2 baris append + header + nomor urut benar.")


if __name__ == "__main__":
    if "--self-check" in sys.argv:
        _self_check()
    else:
        main()
